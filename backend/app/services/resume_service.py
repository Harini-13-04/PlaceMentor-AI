import uuid
import io
import re
import json
import pypdf
import docx
from datetime import datetime, timezone
from typing import List, Optional, Dict, Any
from fastapi import HTTPException, status
from app.database.mongodb import resumes_collection
from app.schemas.resume import (
    ResumeData,
    ResumeCreateRequest,
    ResumeUpdateRequest,
    ResumeSummary,
    ResumeDashboardResponse,
)


def calculate_completion_percentage(doc: dict) -> int:
    score = 0
    
    # Personal Info (15%)
    p_info = doc.get("personal_info") or {}
    if isinstance(p_info, dict):
        if p_info.get("full_name"):
            score += 5
        if p_info.get("email"):
            score += 5
        if p_info.get("phone"):
            score += 5

    # Summary (10%)
    if doc.get("summary"):
        score += 10

    # Education (20%)
    if doc.get("education") and len(doc["education"]) > 0:
        score += 20

    # Experience (20%)
    if doc.get("experience") and len(doc["experience"]) > 0:
        score += 20

    # Skills (15%)
    if doc.get("skills") and len(doc["skills"]) > 0:
        score += 15

    # Projects (10%)
    if doc.get("projects") and len(doc["projects"]) > 0:
        score += 10

    # Additional sections (10%)
    add_count = 0
    if doc.get("certifications") and len(doc["certifications"]) > 0:
        add_count += 1
    if doc.get("achievements") and len(doc["achievements"]) > 0:
        add_count += 1
    if doc.get("languages") and len(doc["languages"]) > 0:
        add_count += 1
    if doc.get("links") and len(doc["links"]) > 0:
        add_count += 1
    if doc.get("custom_sections") and len(doc["custom_sections"]) > 0:
        add_count += 1

    score += min(10, add_count * 3)

    return min(100, score)


async def create_resume(user_id: str, request: ResumeCreateRequest) -> ResumeData:
    now_iso = datetime.now(timezone.utc).isoformat()
    resume_obj = ResumeData(
        id=str(uuid.uuid4()),
        user_id=user_id,
        name=request.name.strip(),
        target_role=request.target_role.strip(),
        experience_level=request.experience_level,
        template=request.template,
        created_at=now_iso,
        updated_at=now_iso,
        completion_percentage=0,
        version=1,
    )

    doc = resume_obj.model_dump()
    doc["completion_percentage"] = calculate_completion_percentage(doc)

    await resumes_collection.insert_one(doc)
    doc.pop("_id", None)
    return ResumeData(**doc)


async def get_user_resumes(user_id: str) -> List[ResumeSummary]:
    cursor = resumes_collection.find({"user_id": user_id}, {"_id": 0})
    resumes = await cursor.to_list(1000)

    resumes.sort(key=lambda r: r.get("updated_at", ""), reverse=True)

    summaries = []
    for r in resumes:
        summaries.append(
            ResumeSummary(
                id=r["id"],
                user_id=r["user_id"],
                name=r.get("name", "Untitled Resume"),
                target_role=r.get("target_role", "Full Stack Developer"),
                experience_level=r.get("experience_level", "Fresher"),
                template=r.get("template", "modern"),
                completion_percentage=r.get("completion_percentage", 0),
                version=r.get("version", 1),
                updated_at=r.get("updated_at", ""),
                created_at=r.get("created_at", ""),
            )
        )
    return summaries


async def get_resume_by_id(resume_id: str, user_id: str) -> Optional[ResumeData]:
    doc = await resumes_collection.find_one(
        {"id": resume_id, "user_id": user_id},
        {"_id": 0}
    )
    if not doc:
        return None
    return ResumeData(**doc)


async def update_resume(resume_id: str, user_id: str, request: ResumeUpdateRequest) -> Optional[ResumeData]:
    existing = await resumes_collection.find_one(
        {"id": resume_id, "user_id": user_id},
        {"_id": 0}
    )
    if not existing:
        return None

    update_dict = request.model_dump(exclude_unset=True)
    if not update_dict:
        return ResumeData(**existing)

    now_iso = datetime.now(timezone.utc).isoformat()
    existing.update(update_dict)
    existing["updated_at"] = now_iso
    existing["version"] = existing.get("version", 1) + 1
    existing["completion_percentage"] = calculate_completion_percentage(existing)

    await resumes_collection.update_one(
        {"id": resume_id, "user_id": user_id},
        {"$set": existing}
    )

    updated_doc = await resumes_collection.find_one(
        {"id": resume_id, "user_id": user_id},
        {"_id": 0}
    )
    return ResumeData(**updated_doc)


async def delete_resume(resume_id: str, user_id: str) -> bool:
    result = await resumes_collection.delete_one({"id": resume_id, "user_id": user_id})
    return result.deleted_count > 0


async def duplicate_resume(resume_id: str, user_id: str) -> Optional[ResumeData]:
    original = await resumes_collection.find_one(
        {"id": resume_id, "user_id": user_id},
        {"_id": 0}
    )
    if not original:
        return None

    now_iso = datetime.now(timezone.utc).isoformat()
    new_doc = dict(original)
    new_doc["id"] = str(uuid.uuid4())
    new_doc["name"] = f"{original.get('name', 'Resume')} (Copy)"
    new_doc["created_at"] = now_iso
    new_doc["updated_at"] = now_iso
    new_doc["version"] = 1
    new_doc["completion_percentage"] = calculate_completion_percentage(new_doc)

    await resumes_collection.insert_one(new_doc)
    new_doc.pop("_id", None)
    return ResumeData(**new_doc)


async def get_dashboard_data(user_id: str) -> ResumeDashboardResponse:
    summaries = await get_user_resumes(user_id)
    total_resumes = len(summaries)
    
    if total_resumes > 0:
        avg_comp = int(sum(s.completion_percentage for s in summaries) / total_resumes)
    else:
        avg_comp = 0

    return ResumeDashboardResponse(
        total_resumes=total_resumes,
        average_completion=avg_comp,
        average_ats_score=None,
        average_defensibility=None,
        resumes=summaries,
    )


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    ext = filename.lower().split('.')[-1]
    if ext == 'pdf':
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_pages = [page.extract_text() or "" for page in reader.pages]
            full_text = "\n".join(text_pages).strip()
            if not full_text:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Could not extract text from PDF file. The document may be scanned, image-only, or empty."
                )
            return full_text
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to process PDF file: {str(e)}"
            )
    elif ext == 'docx':
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            full_text = "\n".join(paragraphs).strip()
            if not full_text:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Could not extract text from DOCX document. File appears to be empty."
                )
            return full_text
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to process DOCX file: {str(e)}"
            )
    elif ext == 'txt':
        try:
            full_text = file_bytes.decode('utf-8', errors='ignore').strip()
            if not full_text:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Text file is empty."
                )
            return full_text
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read text file: {str(e)}"
            )
    elif ext == 'json':
        try:
            full_text = file_bytes.decode('utf-8', errors='ignore').strip()
            if not full_text:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="JSON file is empty."
                )
            return full_text
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read JSON file: {str(e)}"
            )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '.{ext}'. Supported formats are: PDF, DOCX, TXT, JSON."
        )


def parse_resume_text_to_schema(text: str, filename: str) -> dict:
    ext = filename.lower().split('.')[-1]
    if ext == 'json':
        try:
            parsed = json.loads(text)
            if isinstance(parsed, dict):
                return parsed
        except Exception:
            pass

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    full_name = lines[0] if lines else filename.replace('.' + ext, '').replace('_', ' ')

    email = ""
    email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', text)
    if email_match:
        email = email_match.group(0)

    phone = ""
    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
    if phone_match:
        phone = phone_match.group(0)

    linkedin = ""
    linkedin_match = re.search(r'(https?://)?(www\.)?linkedin\.com/in/[\w\.-]+', text, re.IGNORECASE)
    if linkedin_match:
        linkedin = linkedin_match.group(0)

    github = ""
    github_match = re.search(r'(https?://)?(www\.)?github\.com/[\w\.-]+', text, re.IGNORECASE)
    if github_match:
        github = github_match.group(0)

    sections: Dict[str, List[str]] = {
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "projects": [],
        "certifications": [],
        "achievements": [],
        "languages": [],
        "links": [],
        "custom_sections": []
    }

    sec_keywords = {
        "summary": ["summary", "professional summary", "profile", "executive summary", "about me", "objective", "career objective", "overview"],
        "experience": ["experience", "work experience", "professional experience", "employment", "employment history", "work history", "career history", "career highlights", "internships", "relevant experience"],
        "education": ["education", "academic background", "educational background", "academic qualifications", "education & qualifications", "qualifications", "academic"],
        "skills": ["skills", "technical skills", "core competencies", "competencies", "technologies", "skills & tools", "technical proficiencies", "areas of expertise", "expertise"],
        "projects": ["projects", "key projects", "personal projects", "academic projects", "selected projects", "project experience", "portfolio"],
        "certifications": ["certifications", "licenses", "certificates", "certifications & licenses", "licenses & certifications", "credentials"],
        "achievements": ["achievements", "honors & awards", "awards", "key achievements", "honors", "accomplishments"],
        "languages": ["languages", "languages spoken", "language skills"],
        "links": ["links", "websites", "social links", "online presence"],
        "custom_sections": ["publications", "volunteer", "volunteer work", "community involvement", "extracurricular", "organizations", "interests", "hobbies", "additional information"]
    }

    current_sec = "summary"
    for line in lines:
        line_clean = re.sub(r'[^a-zA-Z\s&]', '', line).strip().lower()
        if line_clean.endswith(':'):
            line_clean = line_clean[:-1].strip()

        matched_sec = None
        for sec_name, kw_list in sec_keywords.items():
            if line_clean in kw_list or (len(line_clean) < 35 and any(line_clean == kw or line_clean.startswith(kw) or kw in line_clean for kw in kw_list)):
                matched_sec = sec_name
                break
        if matched_sec:
            current_sec = matched_sec
        else:
            sections[current_sec].append(line)

    summary_text = "\n".join(sections["summary"]).strip() if sections["summary"] else text[:400]

    # EXPERIENCE PARSING - Preserve ALL experience entries
    experience_list = []
    if sections["experience"]:
        cur_exp = None
        for line in sections["experience"]:
            is_header = not line.startswith(('-', '•', '*', '>')) and len(line) < 120 and (
                any(k in line.lower() for k in ["inc", "ltd", "corp", "technologies", "solutions", "university", "lab", "at ", " - ", " – "])
                or re.search(r'\b(20\d{2}|19\d{2}|present|current)\b', line, re.IGNORECASE)
                or cur_exp is None
            )
            if is_header and cur_exp and (cur_exp.get("role") or cur_exp.get("bullets") or cur_exp.get("description")):
                experience_list.append(cur_exp)
                cur_exp = None

            if cur_exp is None:
                role_extracted = ""
                company_extracted = ""
                if " at " in line:
                    parts = line.split(" at ", 1)
                    role_extracted, company_extracted = parts[0].strip(), parts[1].strip()
                elif " - " in line or " – " in line:
                    delim = " - " if " - " in line else " – "
                    parts = line.split(delim, 1)
                    role_extracted, company_extracted = parts[0].strip(), parts[1].strip()
                else:
                    role_extracted = line.strip()

                cur_exp = {
                    "id": str(uuid.uuid4()),
                    "company": company_extracted,
                    "role": role_extracted,
                    "location": "",
                    "start_date": "",
                    "end_date": "",
                    "current": False,
                    "description": line,
                    "bullets": []
                }
            else:
                if line.startswith(('-', '•', '*', '>')):
                    cur_exp["bullets"].append(line.strip('-•* >'))
                else:
                    if cur_exp["description"]:
                        cur_exp["description"] += "\n" + line
                    else:
                        cur_exp["description"] = line
        if cur_exp:
            experience_list.append(cur_exp)

    # EDUCATION PARSING - Preserve ALL education entries
    education_list = []
    if sections["education"]:
        cur_edu = None
        for line in sections["education"]:
            if not cur_edu or (not line.startswith(('-', '•', '*', '>')) and len(line) < 100 and any(k in line.lower() for k in ["university", "college", "institute", "school", "bachelor", "master", "b.e", "b.tech", "m.tech", "bs", "ms", "phd", "degree", "diploma"])):
                if cur_edu:
                    education_list.append(cur_edu)
                inst = line.strip()
                g_match = re.search(r'\b(\d+(\.\d+)?\s*%|\d+(\.\d+)?\s*cgpa|\d+(\.\d+)?\s*gpa|(gpa|cgpa|grade|marks|score)\s*:?\s*\d+(\.\d+)?(%|\s*cgpa|\s*gpa)?)\b', line, re.IGNORECASE)
                ext_grade = g_match.group(0) if g_match else ""
                cur_edu = {
                    "id": str(uuid.uuid4()),
                    "institution": inst,
                    "degree": "",
                    "field_of_study": "",
                    "start_date": "",
                    "end_date": "",
                    "current": False,
                    "gpa": ext_grade,
                    "grade": ext_grade,
                    "description": line
                }
            else:
                if cur_edu:
                    if not cur_edu["degree"] and any(k in line.lower() for k in ["bachelor", "master", "b.e", "b.tech", "m.tech", "bs", "ms", "phd", "degree", "diploma"]):
                        cur_edu["degree"] = line.strip()
                    if not cur_edu["grade"]:
                        g_match = re.search(r'\b(\d+(\.\d+)?\s*%|\d+(\.\d+)?\s*cgpa|\d+(\.\d+)?\s*gpa|(gpa|cgpa|grade|marks|score)\s*:?\s*\d+(\.\d+)?(%|\s*cgpa|\s*gpa)?)\b', line, re.IGNORECASE)
                        if g_match:
                            cur_edu["grade"] = g_match.group(0)
                            cur_edu["gpa"] = g_match.group(0)
                    if cur_edu["description"]:
                        cur_edu["description"] += " | " + line.strip()
                    else:
                        cur_edu["description"] = line.strip()
        if cur_edu:
            education_list.append(cur_edu)

    # SKILLS PARSING - Preserve ALL skill entries
    skills_list = []
    if sections["skills"]:
        extracted_skills = []
        for line in sections["skills"]:
            items = re.split(r'[,|•;]', line)
            for item in items:
                cleaned = item.strip()
                if cleaned and len(cleaned) < 50:
                    extracted_skills.append(cleaned)
        if extracted_skills:
            skills_list.append({
                "id": str(uuid.uuid4()),
                "category": "Technical Skills",
                "skills": extracted_skills
            })

    # PROJECTS PARSING - Preserve ALL project entries
    projects_list = []
    if sections["projects"]:
        cur_proj = None
        for line in sections["projects"]:
            is_proj_header = not line.startswith(('-', '•', '*', '>')) and len(line) < 100 and (cur_proj is None or len(cur_proj.get("bullets", [])) > 0 or len(cur_proj.get("description", "")) > 100)
            if is_proj_header:
                if cur_proj:
                    projects_list.append(cur_proj)
                cur_proj = {
                    "id": str(uuid.uuid4()),
                    "name": line.strip(),
                    "description": line.strip(),
                    "technologies": [],
                    "github_url": "",
                    "live_url": "",
                    "bullets": []
                }
            else:
                if cur_proj:
                    if line.startswith(('-', '•', '*', '>')):
                        cur_proj["bullets"].append(line.strip('-•* >'))
                    else:
                        if cur_proj["description"]:
                            cur_proj["description"] += "\n" + line
                        else:
                            cur_proj["description"] = line
        if cur_proj:
            projects_list.append(cur_proj)

    # CERTIFICATIONS PARSING
    certifications_list = []
    if sections["certifications"]:
        for line in sections["certifications"]:
            if line.strip():
                certifications_list.append({
                    "id": str(uuid.uuid4()),
                    "name": line.strip('-•* '),
                    "issuer": "",
                    "date": "",
                    "url": ""
                })

    # ACHIEVEMENTS PARSING
    achievements_list = []
    if sections["achievements"]:
        for line in sections["achievements"]:
            if line.strip():
                achievements_list.append({
                    "id": str(uuid.uuid4()),
                    "title": line.strip('-•* '),
                    "description": line.strip('-•* '),
                    "date": ""
                })

    # LANGUAGES PARSING
    languages_list = []
    if sections["languages"]:
        for line in sections["languages"]:
            items = re.split(r'[,|•;]', line)
            for item in items:
                if item.strip():
                    languages_list.append({
                        "id": str(uuid.uuid4()),
                        "name": item.strip(),
                        "proficiency": "Fluent"
                    })

    # LINKS PARSING
    links_list = []
    if sections["links"]:
        for line in sections["links"]:
            urls = re.findall(r'https?://[^\s]+', line)
            for url in urls:
                links_list.append({
                    "id": str(uuid.uuid4()),
                    "label": "Link",
                    "url": url
                })

    # CUSTOM SECTIONS PARSING
    custom_sections_list = []
    if sections["custom_sections"]:
        custom_sections_list.append({
            "id": str(uuid.uuid4()),
            "title": "Additional Information",
            "content": "\n".join(sections["custom_sections"]),
            "bullets": [b.strip('-•* ') for b in sections["custom_sections"] if b.startswith(('-', '•', '*', '>'))]
        })

    return {
        "name": filename.replace('.' + ext, '').replace('_', ' ').title(),
        "target_role": "Software Engineer",
        "experience_level": "Mid Level" if experience_list else "Fresher",
        "template": "modern",
        "personal_info": {
            "full_name": full_name,
            "email": email,
            "phone": phone,
            "location": "",
            "linkedin": linkedin,
            "github": github,
            "portfolio": ""
        },
        "summary": summary_text,
        "education": education_list,
        "experience": experience_list,
        "skills": skills_list,
        "projects": projects_list,
        "certifications": certifications_list,
        "achievements": achievements_list,
        "languages": languages_list,
        "links": links_list,
        "custom_sections": custom_sections_list
    }


async def import_resume_file(user_id: str, file_bytes: bytes, filename: str) -> ResumeData:
    raw_text = extract_text_from_file(file_bytes, filename)
    parsed = parse_resume_text_to_schema(raw_text, filename)

    now_iso = datetime.now(timezone.utc).isoformat()
    resume_id = str(uuid.uuid4())

    p_info = parsed.get("personal_info") or {}

    doc = {
        "id": resume_id,
        "user_id": user_id,
        "name": parsed.get("name") or filename,
        "target_role": parsed.get("target_role") or "Software Engineer",
        "experience_level": parsed.get("experience_level") or "Fresher",
        "template": parsed.get("template") or "modern",
        "personal_info": p_info,
        "summary": parsed.get("summary") or "",
        "education": parsed.get("education") or [],
        "experience": parsed.get("experience") or [],
        "skills": parsed.get("skills") or [],
        "projects": parsed.get("projects") or [],
        "certifications": parsed.get("certifications") or [],
        "achievements": parsed.get("achievements") or [],
        "languages": parsed.get("languages") or [],
        "links": parsed.get("links") or [],
        "custom_sections": parsed.get("custom_sections") or [],
        "version": 1,
        "created_at": now_iso,
        "updated_at": now_iso,
    }
    doc["completion_percentage"] = calculate_completion_percentage(doc)

    await resumes_collection.insert_one(doc)
    doc.pop("_id", None)
    return ResumeData(**doc)
