import asyncio
import io
import docx
from fastapi import HTTPException
from app.core.security import create_access_token
from app.database.mongodb import resumes_collection
from app.services.resume_service import (
    parse_resume_text_to_schema,
    import_resume_file,
    get_user_resumes,
    get_resume_by_id,
)


async def run_import_test_suite():
    print("\n==================================================")
    print(" PlaceMentor AI - Resume Import Flow Test Suite ")
    print("==================================================")

    user_a_id = "test-user-import-a"
    user_b_id = "test-user-import-b"

    # Cleanup before run
    await resumes_collection.delete_many({"user_id": {"$in": [user_a_id, user_b_id]}})

    try:
        # Test 1: Import TXT Resume directly via service layer
        print("\n--- [1] import_resume_file (TXT File) ---")
        txt_content = """John Doe
john.doe@example.com | (555) 123-4567 | linkedin.com/in/johndoe

SUMMARY
Experienced Full Stack Engineer specializing in React, Python, and cloud architecture.

EXPERIENCE
Software Engineer - Acme Corp
- Developed scalable microservices handling 1M daily requests.
- Optimized database queries reducing latency by 40%.

EDUCATION
B.S. Computer Science - State University

SKILLS
Python, JavaScript, React, FastAPI, MongoDB, Docker
"""
        txt_bytes = txt_content.encode("utf-8")
        resume_txt = await import_resume_file(user_a_id, txt_bytes, "John_Doe_Resume.txt")

        assert resume_txt.user_id == user_a_id
        assert resume_txt.personal_info.email == "john.doe@example.com"
        assert "(555) 123-4567" in resume_txt.personal_info.phone
        assert "Full Stack Engineer" in resume_txt.summary
        assert resume_txt.completion_percentage > 0
        print(f"[OK] TXT Resume imported & saved: ID={resume_txt.id}, Completion={resume_txt.completion_percentage}%")

        # Verify persisted in user_a resumes list
        user_a_resumes = await get_user_resumes(user_a_id)
        assert any(r.id == resume_txt.id for r in user_a_resumes)
        print("[OK] Verified imported resume appears in get_user_resumes(user_a_id)")

        # Test 2: Import DOCX Resume
        print("\n--- [2] import_resume_file (DOCX File) ---")
        doc = docx.Document()
        doc.add_heading("Jane Smith", 0)
        doc.add_paragraph("jane.smith@example.com | +1 987 654 3210 | github.com/janesmith")
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph("Senior Frontend Architect with 7+ years of React and TypeScript experience.")
        doc.add_heading("Technical Skills", level=1)
        doc.add_paragraph("React, TypeScript, Next.js, Redux, TailwindCSS, Jest")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        resume_docx = await import_resume_file(user_a_id, buffer.getvalue(), "Jane_Smith_Resume.docx")
        assert resume_docx.user_id == user_a_id
        assert resume_docx.personal_info.email == "jane.smith@example.com"
        assert "Frontend Architect" in resume_docx.summary
        print(f"[OK] DOCX Resume imported & saved: ID={resume_docx.id}")

        # Test 3: Import JSON Resume
        print("\n--- [3] import_resume_file (JSON File) ---")
        json_content = """{
            "name": "Alex Taylor Resume",
            "target_role": "Backend Engineer",
            "experience_level": "Senior",
            "personal_info": {
                "full_name": "Alex Taylor",
                "email": "alex.taylor@example.com",
                "phone": "+1 555 999 0000"
            },
            "summary": "Distributed systems specialist.",
            "skills": [{"category": "Languages", "skills": ["Go", "Python", "Rust"]}]
        }"""
        resume_json = await import_resume_file(user_a_id, json_content.encode("utf-8"), "alex_resume.json")
        assert resume_json.name == "Alex Taylor Resume"
        assert resume_json.personal_info.email == "alex.taylor@example.com"
        print(f"[OK] JSON Resume imported & saved: ID={resume_json.id}")

        # Test 4: Unsupported Format Handling
        print("\n--- [4] Unsupported Format Handling ---")
        try:
            await import_resume_file(user_a_id, b"binary content", "malicious.exe")
            assert False, "Should have raised HTTPException"
        except HTTPException as err:
            assert err.status_code == 400
            assert "Unsupported file format" in err.detail
            print("[OK] Unsupported file format correctly raised 400 Bad Request")

        # Test 5: Empty File Handling
        print("\n--- [5] Empty File Handling ---")
        try:
            await import_resume_file(user_a_id, b"", "empty.txt")
            assert False, "Should have raised HTTPException"
        except HTTPException as err:
            assert err.status_code == 400
            assert "empty" in err.detail.lower()
            print("[OK] Empty file correctly raised 400 Bad Request")

        # Test 6: Security & User Isolation Check
        print("\n--- [6] Security & User Isolation Check ---")
        user_b_resumes = await get_user_resumes(user_b_id)
        assert not any(r.id == resume_txt.id for r in user_b_resumes)

        unowned = await get_resume_by_id(resume_txt.id, user_b_id)
        assert unowned is None
        print("[OK] User B cannot access User A's imported resume (returns None / 404)")

        # Test 7: Extraction Safety Test (No Invented Data)
        print("\n--- [7] Extraction Safety Test ---")
        sample_text = """Alice Bob
alice@example.com

SUMMARY
Software Developer

SKILLS
Python, JavaScript
"""
        parsed = parse_resume_text_to_schema(sample_text, "Alice_Resume.txt")
        assert parsed["personal_info"]["full_name"] == "Alice Bob"
        assert parsed["personal_info"]["email"] == "alice@example.com"
        assert parsed["experience"] == []
        assert parsed["education"] == []
        print("[OK] Verified parser extracts only present info without inventing placeholder sections")

        print("\n==================================================")
        print(" ALL RESUME IMPORT TESTS PASSED SUCCESSFULLY! ")
        print("==================================================\n")
    finally:
        await resumes_collection.delete_many({"user_id": {"$in": [user_a_id, user_b_id]}})


if __name__ == "__main__":
    asyncio.run(run_import_test_suite())
